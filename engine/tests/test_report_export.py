"""Tests for the generic CSV/JSON/Excel/PDF Report exporters (Section 14; Phase 7)."""

import csv
import io
import json
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from fios_engine import report_export, reporting
from fios_engine.seed import build_baseline_scenario

AS_OF = date(2026, 7, 15)


@pytest.fixture(scope="module")
def fast_scenario():
    scenario = build_baseline_scenario()
    scenario.monte_carlo_enabled = False
    return scenario


@pytest.fixture(scope="module")
def sample_report(fast_scenario):
    return reporting.assumption_register(fast_scenario, AS_OF)


def test_export_json_round_trips_metadata_and_rows(sample_report):
    payload = json.loads(report_export.export_json(sample_report))

    assert payload["title"] == sample_report.title
    assert payload["metadata"]["scenario_name"] == sample_report.metadata.scenario_name
    assert payload["metadata"]["model_date"] == AS_OF.isoformat()
    assert len(payload["sections"][0]["rows"]) == len(sample_report.sections[0].rows)


def test_export_json_serializes_decimal_values_as_strings(sample_report):
    payload = json.loads(report_export.export_json(sample_report))
    first_row = payload["sections"][0]["rows"][0]

    assert isinstance(first_row["value"], str)


def test_export_csv_contains_metadata_and_section_header(sample_report):
    text = report_export.export_csv(sample_report)
    reader = list(csv.reader(io.StringIO(text)))

    assert reader[0] == ["Report", sample_report.title]
    assert any(row == [sample_report.sections[0].heading] for row in reader)


def test_export_csv_row_count_matches_report_rows(sample_report):
    text = report_export.export_csv(sample_report)
    reader = list(csv.reader(io.StringIO(text)))

    # 5 metadata rows + blank + heading + column header + N data rows.
    data_rows = reader[8:]
    assert len(data_rows) == len(sample_report.sections[0].rows)


def test_export_excel_has_a_metadata_sheet_and_one_sheet_per_section(sample_report):
    workbook = openpyxl.load_workbook(io.BytesIO(report_export.export_excel(sample_report)))

    assert "Metadata" in workbook.sheetnames
    assert len(workbook.sheetnames) == 1 + len(sample_report.sections)


def test_export_excel_converts_decimal_to_a_native_excel_number(fast_scenario):
    """Decimal values become native numeric cells, not strings -- openpyxl itself
    normalizes a whole-number float back to `int` on read, so the meaningful check is
    "not a string, and numerically correct," not "is exactly `float`"."""
    balance_sheet_report = reporting.balance_sheet(fast_scenario, AS_OF)
    assets_section = next(s for s in balance_sheet_report.sections if s.heading == "Assets")
    expected = float(assets_section.rows[0]["balance"])

    workbook = openpyxl.load_workbook(io.BytesIO(report_export.export_excel(balance_sheet_report)))
    assets_sheet = workbook["Assets"]
    header = [cell.value for cell in assets_sheet[1]]
    balance_column = header.index("balance")
    first_data_value = assets_sheet[2][balance_column].value

    assert not isinstance(first_data_value, str)
    assert float(first_data_value) == expected


def test_export_excel_deduplicates_colliding_truncated_sheet_names():
    report = reporting.Report(
        title="Collision Test",
        metadata=reporting.ReportMetadata(AS_OF, "Baseline", "test-version", "n/a"),
        sections=[
            reporting.ReportSection("A" * 40, [{"x": Decimal("1")}]),
            reporting.ReportSection("A" * 39 + "!", [{"x": Decimal("2")}]),
        ],
    )

    workbook = openpyxl.load_workbook(io.BytesIO(report_export.export_excel(report)))

    assert len(set(workbook.sheetnames)) == len(workbook.sheetnames)
    assert len(workbook.sheetnames) == 3


def test_export_pdf_produces_a_non_empty_pdf_document(sample_report):
    pdf_bytes = report_export.export_pdf(sample_report)

    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 0


def test_export_pdf_handles_a_report_with_an_empty_section():
    report = reporting.audit_change_report(build_baseline_scenario(), AS_OF)

    pdf_bytes = report_export.export_pdf(report)

    assert pdf_bytes.startswith(b"%PDF")
